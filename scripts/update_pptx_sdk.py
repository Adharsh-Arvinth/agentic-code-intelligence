"""Update presentation slides with Python SDK details and 50/50 test counts."""

import sys
import os
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor


def update_presentation(path: str):
    print(f"Opening presentation: {path}")
    prs = Presentation(path)

    # 1. Slide 2: Update Enterprise Readiness
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Enterprise Readiness:" in p.text:
                    p.text = (
                        "• Enterprise Readiness: Version-aware index management, 51,234-vector persistent cache, "
                        "Python SDK (agentic_code_intelligence v1.0.0 in dist/), MTEB 2.x compliance, and 50/50 unit tests (100%)."
                    )

    # 2. Slide 12: Update Tools and Tech Stack
    slide12 = prs.slides[11]
    for shape in slide12.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Official Evaluation" in p.text:
                    p.text = (
                        "• Official Evaluation, Indexing & SDK: MTEB v2.21.8 compliance, FAISS IndexFlatIP, "
                        "50 automated pytest unit tests (100%), and Developer Python SDK (agentic_code_intelligence v1.0.0 wheel in dist/)"
                    )

    # 3. Slide 16: Update Brownie points
    slide16 = prs.slides[15]
    for shape in slide16.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "4. Native MTEB 2.x SearchProtocol" in p.text or "4. Native MTEB 2.x Protocol" in p.text:
                    p.text = "4. Native MTEB 2.x Protocol + Python SDK + 50/50 Unit Tests"
                elif "Implements both MTEB 2.x EncoderProtocol" in p.text or "Implements MTEB 2.x SearchProtocol" in p.text:
                    p.text = (
                        "Implements MTEB 2.x SearchProtocol, ships prebuilt Python SDK wheel (agentic_code_intelligence v1.0.0), "
                        "supports multi-version index isolation, and passes 50/50 automated pytest unit tests (100%)."
                    )

    # 4. Slide 17: Update Checklist with SDK bullet and Demo.mp4
    slide17 = prs.slides[16]
    for shape in slide17.shapes:
        if shape.has_text_frame:
            paras = shape.text_frame.paragraphs
            if any("Working prototype code" in p.text for p in paras):
                for p in paras:
                    if "Demo video:" in p.text:
                        p.text = (
                            "• Demo video: Demo.mp4 (in repository, 20MB) & Google Drive: "
                            "https://drive.google.com/file/d/1JQQXUEOY81HSYag60_mDM9-FCE1qdSzh/view?usp=sharing"
                        )
                # Check if SDK bullet already exists
                has_sdk = any("APK / SDK" in p.text for p in paras)
                if not has_sdk:
                    p_sdk = shape.text_frame.add_paragraph()
                    p_sdk.text = "• APK / SDK (Y/N): YES (Production Python SDK 'agentic_code_intelligence' v1.0.0 wheel in dist/ + examples/sdk_quickstart.py)"

    prs.save(path)
    print(f"Successfully saved updated presentation to: {path}")


if __name__ == "__main__":
    target_pptx = r"c:\Users\adhar\Desktop\samsung\Samsung_PRISM_Generative_AI_Hackathon_Presentation_Updated.pptx"
    update_presentation(target_pptx)
